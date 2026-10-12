"""What an uploaded file is, worked out before it joins the library.

Everything here is plain functions over a file or over text already read from it:
the format (from the file's content, not just its name), the details the file
states about itself (its properties, an EPUB's package, the copyright page), and
the single "type" the review screen offers (novel, short story, paper, manual ...)
that maps onto the library's kind, book form and document type.

Nothing here touches the library; app.py stages the file, runs these, and keeps the
result until the reader saves or discards it.
"""

from __future__ import annotations

import datetime
import html
import re
import zipfile
from pathlib import Path
from typing import Iterable, List, Optional

import doc_kind
import doc_types
import shelves

FORMAT_LABELS = {
    "pdf": "PDF",
    "epub": "EPUB",
    "word": "Word document",
    "excel": "Excel workbook",
    "powerpoint": "PowerPoint deck",
    "text": "Plain text",
    "markdown": "Markdown",
    "csv": "CSV table",
}
SUFFIXES = {
    "pdf": ".pdf", "epub": ".epub", "word": ".docx", "excel": ".xlsx",
    "powerpoint": ".pptx", "text": ".txt", "markdown": ".md", "csv": ".csv",
}
_EDGE_PAGES = 15  # the front matter and the last pages hold the copyright and imprint
_FIRST_YEAR = 1450  # nothing printed earlier belongs in a personal library


# ------------------------------------------------------------- the format --
def sniff_format(path: str | Path) -> Optional[str]:
    """The file type its content shows (a PDF's header, an Office or EPUB
    package's parts); None for plain text, which has no signature, or when the
    content is not recognised."""
    try:
        with open(path, "rb") as handle:
            head = handle.read(1024)
    except OSError:
        return None
    if b"%PDF-" in head:
        return "pdf"
    if not head.startswith(b"PK\x03\x04"):
        return None
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "mimetype" in names and archive.read("mimetype")[:40].strip() == b"application/epub+zip":
                return "epub"
            if "word/document.xml" in names:
                return "word"
            if "xl/workbook.xml" in names:
                return "excel"
            if "ppt/presentation.xml" in names:
                return "powerpoint"
            if "META-INF/container.xml" in names:
                return "epub"
    except (zipfile.BadZipFile, OSError, KeyError, RuntimeError):
        return None
    return None


def resolve_format(named_type: str, path: str | Path) -> tuple[str, Optional[str]]:
    """(the type to read the file as, a note for the reader or None). The
    content wins over the extension, so a PDF saved as ".docx" is read as a PDF."""
    sniffed = sniff_format(path)
    if sniffed and sniffed != named_type:
        return sniffed, f"The name says {FORMAT_LABELS.get(named_type, named_type)}, but the content is {FORMAT_LABELS[sniffed]}; it will be read as that."
    if named_type in {"pdf", "epub", "word", "excel", "powerpoint"} and sniffed is None:
        return named_type, f"This does not look like a valid {FORMAT_LABELS[named_type]}; it may fail to open."
    if named_type in {"text", "markdown", "csv"}:
        try:
            with open(path, "rb") as handle:
                if b"\x00" in handle.read(4096):
                    return named_type, "This looks like a binary file, not text."
        except OSError:
            pass
    return named_type, None


# ------------------------------------------------- what the file says itself --
_JUNK_TITLE = re.compile(
    r"^(?:microsoft (?:word|excel|powerpoint)\b.*|untitled\b.*|document\d*|layout \d+|slide \d+|new document.*|"
    r".*\.(?:docx?|indd|qxd|pdf|tex|xlsx?|pptx?|rtf|txt)|\W*)$",
    re.IGNORECASE,
)
_JUNK_STEM = re.compile(
    r"^(?:[\W\d_]*|[0-9a-f]{16,}|(?:scan(?:ned)?|doc|document|img|image|file|download|untitled|new[ _-]?document|page|slide)[\W\d_]*)$",
    re.IGNORECASE,
)


def plausible_title(value: Optional[str]) -> Optional[str]:
    """A title a file's properties give, unless it is the authoring tool's
    placeholder ("Microsoft Word - report.docx")."""
    text = " ".join((value or "").split())
    if len(text) < 2 or len(text) > 300 or _JUNK_TITLE.match(text):
        return None
    return text


def junk_stem(stem: str) -> bool:
    """A file name that says nothing about the book ("scan0012", a hash)."""
    return bool(_JUNK_STEM.match(stem.strip()))


def _tag(opf: str, name: str) -> List[str]:
    found = re.findall(rf"<(?:\w+:)?{name}\b[^>]*>(.*?)</(?:\w+:)?{name}>", opf, re.IGNORECASE | re.DOTALL)
    return [" ".join(html.unescape(re.sub(r"<[^>]+>", "", item)).split()) for item in found if item.strip()]


def _epub_package(path: str | Path) -> str:
    """The text of an EPUB's package document (its OPF), read as text so no XML
    from an untrusted file is ever parsed."""
    with zipfile.ZipFile(path) as archive:
        container = archive.read("META-INF/container.xml")[:200_000].decode("utf-8", "replace")
        match = re.search(r'full-path\s*=\s*["\']([^"\']+)["\']', container)
        if not match:
            return ""
        return archive.read(match.group(1))[:2_000_000].decode("utf-8", "replace")


def valid_isbn(value: str) -> Optional[str]:
    """The ISBN with hyphens and spaces removed, or None when it is not one."""
    digits = re.sub(r"[-\s]", "", value or "").upper()
    if re.fullmatch(r"\d{13}", digits):
        return digits if sum((1 if i % 2 == 0 else 3) * int(c) for i, c in enumerate(digits)) % 10 == 0 else None
    if re.fullmatch(r"\d{9}[\dX]", digits):
        values = [int(c) for c in digits[:9]] + [10 if digits[-1] == "X" else int(digits[-1])]
        return digits if sum((10 - i) * v for i, v in enumerate(values)) % 11 == 0 else None
    return None


def _year_of(value: str) -> Optional[int]:
    match = re.search(r"\b(1[4-9]\d{2}|20\d{2})\b", value or "")
    return _plausible_year(int(match.group(1))) if match else None


def _plausible_year(year: int) -> Optional[int]:
    return year if _FIRST_YEAR <= year <= datetime.date.today().year + 1 else None


def embedded_details(path: str | Path, file_type: str) -> dict:
    """What a file states about itself. `author` is the raw value (it may be a
    translator or an institution; shelves.normalize_author judges it). Only keys the
    file actually has; {} when it cannot be read."""
    out: dict = {}
    try:
        if file_type == "epub":
            opf = _epub_package(path)
            if title := plausible_title(next(iter(_tag(opf, "title")), None)):
                out["title"] = title
            if creators := _tag(opf, "creator"):
                out["author"] = ", ".join(dict.fromkeys(creators))
            if publisher := next(iter(_tag(opf, "publisher")), None):
                out["publisher"] = publisher[:300]
            if language := next(iter(_tag(opf, "language")), None):
                if re.fullmatch(r"[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})*", language):
                    out["language"] = language.lower()
            if year := _year_of(next(iter(_tag(opf, "date")), "")):
                out["published_year"] = year
            for identifier in _tag(opf, "identifier"):
                if isbn := valid_isbn(re.sub(r"^urn:isbn:", "", identifier, flags=re.IGNORECASE)):
                    out["isbn_13" if len(isbn) == 13 else "isbn_10"] = isbn
            if description := next(iter(_tag(opf, "description")), None):
                out["description"] = description[:2000]
            if subjects := _tag(opf, "subject"):
                out["genres"] = list(dict.fromkeys(subjects))[:8]
        elif file_type == "pdf":
            import fitz

            with fitz.open(str(path)) as document:
                meta = document.metadata or {}
            if title := plausible_title(meta.get("title")):
                out["title"] = title
            if meta.get("author"):
                out["author"] = meta["author"]
            if meta.get("subject"):
                out["description"] = " ".join(meta["subject"].split())[:2000]
        elif file_type in {"word", "powerpoint"}:
            if file_type == "word":
                from docx import Document

                props = Document(str(path)).core_properties
            else:
                from pptx import Presentation

                props = Presentation(str(path)).core_properties
            if title := plausible_title(props.title):
                out["title"] = title
            if props.author:
                out["author"] = props.author
        elif file_type == "excel":
            from openpyxl import load_workbook

            workbook = load_workbook(filename=str(path), read_only=True)
            try:
                if title := plausible_title(workbook.properties.title):
                    out["title"] = title
                if workbook.properties.creator:
                    out["author"] = workbook.properties.creator
            finally:
                workbook.close()
    except Exception:
        return out
    return out


# ------------------------------------------------------ the copyright page --
_ISBN_TEXT = re.compile(r"\bISBN(?:-1[03])?[:\s]*((?:97[89][-\s]?)?\d[\d\-\s]{7,15}[\dXx])\b", re.IGNORECASE)
_YEAR = r"((?:1[4-9]|20)\d{2})(?!\d|[.,]\d)"  # a bare year, not the start of "1617.56"
_COPYRIGHT_YEAR = re.compile(
    rf"(?:(?:©|\bcopyright\b)[^\n\d]{{0,40}}|\(c\)\s*){_YEAR}(?:\s*[,&–-]\s*{_YEAR})*", re.IGNORECASE
)
_FIRST_PUBLISHED = re.compile(r"\bfirst\s+(?:published|printed|edition)\b[^\n\d]{0,40}((?:1[4-9]|20)\d{2})", re.IGNORECASE)
_PUBLISHER = re.compile(
    r"(?:\bpublished\s+(?:in\s+[A-Za-z ,]{3,40}?\s+)?by|^\s*publisher\s*:)\s+((?-i:[A-Z])[^\n,;©(]{2,70})",
    re.IGNORECASE | re.MULTILINE,
)
_RENEWAL = re.compile(r"\b(?:renew\w*|reprint\w*|reissue\w*|revised)\b[^\n\d]{0,20}(?:1[4-9]|20)\d{2}", re.IGNORECASE)
_ADDRESS = re.compile(r"\s+\d{1,5}\s+\S.*$")
# What follows a publisher's name on the same line of a copyright page.
_AFTER_PUBLISHER = re.compile(r"\b(?:all rights|printed|first|copyright|isbn|typeset|in association|,?\s+an imprint)\b.*$", re.IGNORECASE)


def front_matter_details(page_texts: Iterable[str]) -> dict:
    """The ISBN, the year and the publisher printed on a book's copyright page. Only
    the opening and closing pages are read. The year is the latest one printed
    there, which is this edition's own: earlier ones belong to stories reprinted
    in an anthology or to a first edition, and a renewal date is not counted."""
    pages = [text or "" for text in page_texts]
    edges = "\n".join(pages[:_EDGE_PAGES] + pages[-4:])
    dated = _RENEWAL.sub(" ", edges)
    out: dict = {}
    for match in _ISBN_TEXT.finditer(edges):
        isbn = valid_isbn(match.group(1))
        if isbn:
            out.setdefault("isbn_13" if len(isbn) == 13 else "isbn_10", isbn)
    years = [int(year) for year in re.findall(r"((?:1[4-9]|20)\d{2})", " ".join(m.group(0) for m in _COPYRIGHT_YEAR.finditer(dated)))]
    years += [int(m.group(1)) for m in _FIRST_PUBLISHED.finditer(dated)]
    years = [year for year in years if _plausible_year(year)]
    if years:
        out["published_year"] = max(years)
    publisher = _PUBLISHER.search(edges)
    if publisher:
        name = _ADDRESS.sub("", _AFTER_PUBLISHER.sub("", re.sub(r"\s+", " ", publisher.group(1)))).strip(" .,:;-")
        if 3 <= len(name) <= 80:
            out["publisher"] = name
    return out


# --------------------------------------------------------- the type menu --
# One menu for what the reader thinks of as "what is this": a book's form, or
# the kind of document. Each maps onto the library's kind, form and doc type.
BOOK_TYPES = ("novel", "novella", "short_story", "collection", "nonfiction", "book")
DOCUMENT_TYPES = tuple(doc_types.TYPES)


def type_options() -> List[dict]:
    """The type menu, each entry with the top-level shelf that type belongs on."""
    books = [{"value": value, "label": label, "group": "Books", "shelf": shelves.TOP_LEVEL["book"]} for value, label in (
        ("novel", "Novel"), ("novella", "Novella"), ("short_story", "Short story"), ("collection", "Collection"),
        ("nonfiction", "Non-fiction book"), ("book", "Book (form not decided)"),
    )]
    documents = [
        {"value": value, "label": "Other document" if value == "other" else doc_types.LABELS[value], "group": "Documents",
         "shelf": shelves.DOCUMENT_TOP_LEVEL[value]}
        for value in DOCUMENT_TYPES
    ]
    return books + documents


def type_patch(value: str) -> dict:
    """The kind, form and document type to store for a chosen type."""
    if value in BOOK_TYPES:
        return {"kind": "book", "form": "auto" if value == "book" else value, "doc_type": "auto"}
    if value == "paper":
        return {"kind": "paper", "form": "auto", "doc_type": "paper"}
    if value in DOCUMENT_TYPES:
        return {"kind": "document", "form": "auto", "doc_type": value}
    raise ValueError(f"unknown type: {value}")


def type_of(kind: Optional[str], form: Optional[str], doc_type: Optional[str]) -> str:
    """The menu entry for what the rules found."""
    if kind == "book":
        return form if form in BOOK_TYPES else "book"
    if kind == "paper":
        return "paper"
    return doc_type if doc_type in DOCUMENT_TYPES else "other"
