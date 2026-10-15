#!/usr/bin/env python3
"""Synthetic benchmark: documents with planted facts, and questions whose
answers are known by construction.

    python eval/synth_corpus.py ../library/synth          # writes ~60 files, the
                                                          # query set and the answer set

Why this exists. The book suites are judged by looking at what the system
returned and keeping the good hits, so the system partly grades itself, and
tuning against them (the relevance gate, the reranker) makes every number
in-sample. Here the order is reversed: facts are *generated first* (a manager's
name, a pallet count, a charter year), written into documents of seven kinds
(docx, pdf, pptx, xlsx, csv, md, txt), and only then turned into questions. The
right answer for each question is known before any model has seen it.

What it deliberately makes hard:

* **Siblings.** Fourteen depot reports share one template, so "who manages the
  Calanford depot" is answered by a sentence that looks like thirteen others.
  Only the name tells them apart.
* **Shared names.** Place names are reused across document kinds (a depot, a
  survey, an office handbook and a town chronicle can all be "Calanford"), so
  the question has to be read, not just keyword-matched.
* **Paraphrase.** Half the questions avoid the words the document used.
* **Reverse lookups.** "Which depot employs 212 people?" -- find the entity
  from a value.
* **Carriers.** Facts sit in prose, a Word table, a spreadsheet row, a slide's
  speaker notes and a PDF page (page numbers are judged too).
* **Comparisons** that need two sources, **absent entities** (a depot that does
  not exist) and **absent attributes** (a real depot, a fact nobody wrote down).

Output: the files, plus ``queries_synth.jsonl`` (retrieval) and
``answers_synth.jsonl`` (Ask mode). Same seed, same text; only the files'
embedded timestamps differ. After writing, every judged phrase is looked up in
the rendered files with plain python-docx / openpyxl / python-pptx / PyMuPDF --
independent of the service's own parsers -- so a broken generator fails here,
not as a mysterious retrieval miss.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import re
import textwrap
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

SEED = 20261014

# --------------------------------------------------------------------- names

_ONSETS = ["Bran", "Cal", "Dar", "Esk", "Fen", "Gal", "Hal", "Iver", "Jor", "Kes", "Lor", "Mar",
           "Nor", "Orm", "Pel", "Quen", "Ros", "Sel", "Tav", "Ulm", "Var", "Wen", "Yor", "Zan",
           "Ast", "Bel", "Cor", "Del", "Eld", "Fal"]
_CODAS = ["ford", "wick", "mere", "ton", "by", "dale", "haven", "stead", "moor", "ley", "ham",
          "burn", "gate", "holm"]
_FIRST = ["Orlin", "Maren", "Teodor", "Ilse", "Corwin", "Sunniva", "Anselm", "Odette", "Radek",
          "Lenka", "Bram", "Iselde", "Tobias", "Marisol", "Kasimir", "Elowen", "Dmitri", "Fiora",
          "Gunnar", "Halima", "Idris", "Jolene", "Konrad", "Linnea", "Matthias", "Nadia", "Osric",
          "Petra", "Quentin", "Rosalind", "Stellan", "Tamsin", "Ulrich", "Vesna", "Wilhelm",
          "Xenia", "Yusuf", "Zora", "Alaric", "Beatrix"]
_LAST_A = ["Hart", "Lind", "Okon", "Petr", "March", "Duqu", "Vass", "Brack", "Sorr", "Tamm", "Wolk",
           "Ander", "Keld", "Rask", "Quill", "Nord", "Fenn", "Gall", "Holm", "Ivar", "Jans", "Korv",
           "Laud", "Moss", "Nash", "Orsk", "Pryn", "Rudd", "Stra", "Thorn"]
_LAST_B = ["ford", "berg", "ski", "son", "mann", "ez", "ova", "ley"]
_FUNDERS = ["Hollis Marine Trust", "Veldt Ocean Fund", "Marrow Bay Foundation", "Tidewell Institute",
            "Corrin Reef Society", "Ashgrove Research Council", "Penhallow Trust", "Sorrel Coastal Fund"]
_PRODUCT_NAMES = ["Velor", "Quanta", "Nimbus", "Tarn", "Orik", "Sable", "Halo", "Ember", "Lumen",
                  "Vanta", "Brisk", "Kestrel"]
_PRODUCT_NOUNS = ["logger", "sensor", "beacon", "probe"]

# Words the "absent attribute" questions ask about. None of them may appear in
# any generated file -- verify() checks -- or the right answer would be
# "refuse" for a question the corpus half-answers.
ABSENT_TOPICS = {
    "depot": ["What is the annual electricity bill of {E}?", "Who is the deputy manager of {E}?",
              "How many forklifts does {E} operate?"],
    "study": ["Which journal published the findings of {E}?", "What was the ethics approval number of {E}?",
              "How many volunteers took part in {E}?"],
    "policy": ["How many paid holiday days do employees get at {E}?", "What is the dress code at {E}?",
               "Who is the head of human resources at {E}?"],
    "product": ["What is the screen resolution of {E}?", "Which colours does {E} come in?",
                "How many units of {E} were sold last year?"],
    "town": ["Who is the mayor of {E}?", "What is the annual rainfall in {E}?",
             "What is the name of the football club in {E}?"],
    "roster": ["What is the postal address of {E}?", "How many desks does {E} have?"],
}
ABSENT_KEYWORDS = ["electricity", "deputy", "forklift", "journal", "ethics", "volunteer", "holiday",
                   "dress code", "human resources", "screen", "colour", "sold", "mayor", "rainfall",
                   "football", "postal", "desks"]

OFF_TOPIC = [
    "how do I repair a bicycle tyre", "recipe for a lemon drizzle cake", "who won the 1998 football world cup",
    "symptoms of the common flu", "how to train for a marathon", "best way to learn the guitar",
    "how does compound interest work", "what is the capital of Mongolia", "how do vaccines work",
    "tips for growing tomatoes at home", "explain the offside rule", "how to write a cover letter",
]

FILLER = {
    "depot": ["Safety walks are held every morning before the first shift begins.",
              "Dock doors are labelled by lane and inspected at the start of each week.",
              "Returned goods are sorted on the mezzanine before they are restocked.",
              "Drivers hand their paperwork to the gatehouse before leaving the yard.",
              "Damaged pallets are tagged in red and set aside for repair.",
              "The yard is swept daily to keep the loading lanes clear."],
    "study": ["Fragments were attached to the substrate with marine epoxy and photographed at regular intervals.",
              "Field teams recorded water temperature and visibility at each visit.",
              "Wave exposure at the plots was moderate throughout the observation period.",
              "Photographs were scored by two observers working independently.",
              "Colonies were checked for signs of disease, predation and sediment stress.",
              "The plots were marked with numbered stakes to make relocation easy."],
    "policy": ["Employees are expected to read the handbook during their first fortnight.",
               "Questions about these rules can be raised with a line manager at any time.",
               "Exceptions are recorded in writing and kept with the personnel file.",
               "The handbook applies to permanent and fixed-term staff alike.",
               "Changes take effect once they have been announced to the whole office.",
               "Printed copies are kept at reception for anyone who prefers paper."],
    "product": ["The housing is moulded from recycled polymer and finished with a matte coating.",
                "A single button on the side starts and stops recording.",
                "Firmware updates are installed over the supplied cable.",
                "The unit mounts on a standard bracket using two screws.",
                "A status light shows green when the device is ready.",
                "Data is stored on removable memory that can be swapped in seconds."],
    "town": ["The market square filled with stalls whenever the harvest was good.",
             "Merchants from the surrounding valleys brought wool, grain and salt to trade.",
             "The church bell marked the hours for the labourers in the fields.",
             "A guild of weavers met in the long hall behind the tavern.",
             "Winters were hard, and the river froze for weeks at a time.",
             "Travellers spoke well of the inns that lined the eastern road.",
             "Disputes over grazing land were settled by elders under the old oak."],
    "roster": ["Figures are provisional until the audit is signed.",
               "Budgets are reviewed by the finance group each autumn.",
               "Sites are listed in the order they joined the programme."],
}

# ---------------------------------------------------------------- attributes
#
# Each attribute: how its value is drawn (kind/lo/hi/unit/prefix/suffix) and
# the wording of its facts and questions. Templates use {E} ("the Calanford
# depot"), {Ec} (capitalised) and {v} (the value, unit included). A fact
# template must put {v} in one piece -- "212 people", never "212 of its people"
# -- because the judged phrase is the formatted value.
#   facts   how the document states it
#   direct  a question reusing the document's own words
#   para    a question that avoids them
#   reverse a question that gives the value and asks for the entity (None when
#           the value is not unique across entities)

def _attr(key, kind, facts, direct, para, reverse=None, **extra):
    return {"key": key, "kind": kind, "facts": facts, "direct": direct, "para": para,
            "reverse": reverse, **extra}


ATTRS = {
    "depot": [
        _attr("manager", "person", ["{Ec} is managed by {v}.", "Day-to-day operations at {E} are run by site manager {v}."],
              ["Who manages {E}?", "Who is the site manager of {E}?"],
              ["Which person is in charge of running {E}?", "Who heads up operations at {E}?"],
              ["Which depot is managed by {v}?", "At which depot does {v} serve as site manager?"]),
        _attr("staff", "int", ["{Ec} employs {v}.", "The workforce at {E} totals {v}."],
              ["How many people does {E} employ?", "How many people work at {E}?"],
              ["What is the headcount at {E}?", "How big is the workforce at {E}?"],
              ["Which depot employs {v}?", "Which depot has a workforce of {v}?"], lo=60, hi=480, unit=" people"),
        _attr("opened", "int", ["{Ec} opened in {v}.", "Operations at {E} began in {v}."],
              ["In what year did {E} open?", "When did {E} open?"],
              ["Since which year has {E} been in service?", "What is the founding year of {E}?"],
              ["Which depot opened in {v}?"], lo=1961, hi=2019, sep=False, cmp=("earlier", "later", "opening year")),
        _attr("ontime", "float", ["On-time delivery at {E} reached {v}.", "{Ec} dispatched {v} of its orders on time."],
              ["What is the on-time delivery rate at {E}?", "What percentage of orders does {E} deliver on time?"],
              ["How punctual are shipments leaving {E}?", "How reliable is {E} at meeting its delivery windows?"],
              ["Which depot reached {v} on-time delivery?"], lo=82.0, hi=99.5, unit=" percent",
              cmp=("higher", "lower", "on-time delivery rate")),
        _attr("pallets", "int", ["{Ec} handles about {v} per week.", "Weekly throughput at {E} is {v}."],
              ["How many pallets does {E} handle each week?", "What is the weekly pallet volume at {E}?"],
              ["How much freight moves through {E} in a week?", "What weekly volume does {E} process?"],
              ["Which depot handles {v} per week?"], lo=800, hi=9000, unit=" pallets",
              cmp=("higher", "lower", "weekly pallet volume")),
        _attr("cost", "float", ["{Ec} charges {v} to handle one pallet."],
              ["What is the cost per pallet at {E}?", "How much does it cost per pallet at {E}?"],
              ["What does {E} charge to move one pallet?", "What is the per-pallet handling price at {E}?"],
              ["Which depot has a cost per pallet of {v}?"], lo=20.0, hi=89.99, prefix="€", carrier="table",
              label="Cost per pallet", cmp=("higher", "lower", "cost per pallet")),
    ],
    "study": [
        _attr("lead", "person", ["{Ec} was led by {v}.", "The principal investigator of {E} was {v}."],
              ["Who led {E}?", "Who was the principal investigator of {E}?"],
              ["Which scientist directed {E}?", "Who was the lead researcher behind {E}?"],
              ["Which survey was led by {v}?"]),
        _attr("funder", "name", ["{Ec} was funded by the {v}.", "Funding for {E} came from the {v}."],
              ["Who funded {E}?", "Which organisation paid for {E}?"],
              ["Where did the money for {E} come from?", "Which body financed {E}?"], None, pool=_FUNDERS),
        _attr("sample", "int", ["{Ec} examined {v}.", "The sample for {E} comprised {v}."],
              ["How many fragments were examined in {E}?", "What was the sample size of {E}?"],
              ["How many pieces of coral did {E} follow?", "What number of fragments was monitored in {E}?"],
              ["Which survey examined {v}?"], lo=120, hi=990, unit=" fragments",
              cmp=("larger", "smaller", "sample size")),
        _attr("duration", "int", ["{Ec} ran for {v}.", "Monitoring at {E} continued for {v}."],
              ["How long did {E} run?", "How many months did monitoring last in {E}?"],
              ["What was the time span of {E}?", "For how long was {E} carried out?"],
              ["Which survey ran for {v}?"], lo=6, hi=36, unit=" months", cmp=("longer", "shorter", "duration")),
        _attr("survival", "float", ["Survival of outplanted fragments in {E} was {v}.", "{Ec} recorded a survival rate of {v}."],
              ["What was the survival rate in {E}?", "What percentage of fragments survived in {E}?"],
              ["How well did the outplanted corals fare in {E}?", "What fraction of the corals lived through {E}?"],
              ["Which survey recorded a survival rate of {v}?"], lo=40.0, hi=95.0, unit=" percent",
              cmp=("higher", "lower", "survival rate")),
    ],
    "policy": [
        _attr("travel", "int", ["Travel advances at {E} are capped at {v}.", "{Ec} reimburses travel up to {v} per trip without further approval."],
              ["What is the travel reimbursement limit at {E}?", "How much can staff at {E} claim for travel without extra approval?"],
              ["Up to what amount can a trip be reimbursed at {E}?", "What is the spending ceiling for journeys at {E}?"],
              ["Which office caps travel advances at {v}?"], lo=150, hi=950, prefix="€", sep=False,
              cmp=("higher", "lower", "travel reimbursement limit")),
        _attr("approver", "person", ["Expenses above the cap at {E} must be approved by {v}.", "{Ec} routes large expense claims to {v}."],
              ["Who approves large expense claims at {E}?", "Who signs off expenses above the limit at {E}?"],
              ["Which person has the final say on big claims at {E}?", "Whose approval do large reimbursements need at {E}?"],
              ["At which office does {v} approve large expense claims?"]),
        _attr("notice", "int", ["Employees at {E} give {v} notice before resigning.", "The notice period at {E} is {v}."],
              ["What is the notice period at {E}?", "How many days notice must employees at {E} give?"],
              ["How long in advance must someone at {E} announce they are leaving?", "How much warning does {E} require before a resignation?"],
              ["Which office has a notice period of {v}?"], lo=30, hi=180, unit=" days", step=5,
              cmp=("longer", "shorter", "notice period")),
        _attr("review", "int", ["The handbook of {E} is reviewed every {v}.", "{Ec} revisits its policies on a cycle of {v}."],
              ["How often is the handbook of {E} reviewed?", "What is the policy review cycle at {E}?"],
              ["How frequently does {E} update its policies?", "After how long are the rules at {E} re-examined?"],
              ["Which office reviews its handbook every {v}?"], lo=6, hi=48, unit=" months",
              cmp=("longer", "shorter", "policy review cycle")),
    ],
    "product": [
        _attr("weight", "int", ["{Ec} weighs {v}.", "The mass of {E} is {v}."],
              ["How much does {E} weigh?", "What is the weight of {E}?"],
              ["How heavy is {E}?", "What does {E} come to on the scales?"],
              ["Which device weighs {v}?"], lo=90, hi=950, unit=" grams", cmp=("heavier", "lighter", "weight")),
        _attr("price", "int", ["{Ec} retails for {v}.", "The list price of {E} is {v}."],
              ["What does {E} cost?", "What is the list price of {E}?"],
              ["How much do I have to pay for {E}?", "What is {E} priced at?"],
              ["Which device retails for {v}?"], lo=49, hi=899, prefix="€", sep=False,
              cmp=("higher", "lower", "list price")),
        _attr("battery", "float", ["{Ec} runs for {v} on a single charge.", "Battery life of {E} is {v}."],
              ["How long does the battery of {E} last?", "What is the battery life of {E}?"],
              ["How many hours can {E} run before it needs charging?", "How long will {E} operate unplugged?"],
              ["Which device lasts {v} on a charge?"], lo=4.0, hi=72.0, unit=" hours",
              cmp=("longer", "shorter", "battery life")),
        _attr("depth", "int", ["{Ec} is waterproof down to {v}.", "Do not submerge {E} below {v}."],
              ["How deep can {E} be submerged?", "What is the maximum depth rating of {E}?"],
              ["How far underwater can {E} go?", "What is the dive limit of {E}?"],
              ["Which device is waterproof down to {v}?"], lo=5, hi=400, unit=" metres", carrier="notes"),
        _attr("warranty", "int", ["{Ec} comes with a warranty of {v}.", "Customers receive {v} of warranty cover on {E}."],
              ["How long is the warranty on {E}?", "What warranty period does {E} have?"],
              ["For how many years is {E} covered if it breaks?", "How long does the manufacturer stand behind {E}?"],
              None, lo=1, hi=10, unit=" years"),
    ],
    "town": [
        _attr("founded", "int", ["{Ec} was founded in {v}.", "Records place the founding of {E} in {v}."],
              ["When was {E} founded?", "In what year was {E} founded?"],
              ["How old is {E}? When was it established?", "What year was {E} first settled?"],
              ["Which town was founded in {v}?"], lo=1050, hi=1299, sep=False, cmp=("earlier", "later", "founding year")),
        _attr("founder", "person", ["{Ec} was founded by {v}.", "Tradition credits {v} with establishing {E}."],
              ["Who founded {E}?", "Who was the founder of {E}?"],
              ["Which person established {E}?", "Who is credited with setting up {E}?"],
              ["Which town was founded by {v}?"]),
        _attr("charter", "int", ["{Ec} received its royal charter in {v}.", "The charter of {E} dates from {v}."],
              ["When did {E} receive its charter?", "What year was the charter of {E} granted?"],
              ["In which year did {E} gain official town rights?", "When was {E} granted its royal privileges?"],
              ["Which town received its charter in {v}?"], lo=1300, hi=1499, sep=False,
              cmp=("earlier", "later", "charter year")),
        _attr("population", "int", ["The census counted {v} in {E}.", "{Ec} had {v} at the last count."],
              ["What is the population of {E}?", "How many inhabitants does {E} have?"],
              ["How many people live in {E}?", "How big is {E} in terms of residents?"],
              ["Which town has {v}?"], lo=1200, hi=48000, unit=" inhabitants", cmp=("larger", "smaller", "population")),
        _attr("bridge", "int", ["The old bridge at {E} spans {v}.", "{Ec} is known for its stone bridge, {v} long."],
              ["How long is the bridge at {E}?", "What is the span of the old bridge at {E}?"],
              ["What is the length of the river crossing at {E}?", "How far does the stone bridge in {E} reach?"],
              ["Which town has a bridge {v} long?"], lo=40, hi=400, unit=" metres", cmp=("longer", "shorter", "bridge")),
    ],
    "roster": [
        _attr("code", "int", ["{Ec} is booked under cost centre {v}."],
              ["What is the cost centre code of {E}?", "Which cost centre does {E} use?"],
              ["Under which accounting code is {E} booked?", "What code identifies the finances of {E}?"],
              ["Which site uses cost centre {v}?"], lo=1000, hi=9999, prefix="CC-", sep=False, carrier="row",
              label="Cost centre"),
        _attr("budget", "int", ["{Ec} has an annual budget of {v}."],
              ["What is the annual budget of {E}?", "How much money is allocated to {E} each year?"],
              ["What yearly funding does {E} receive?", "What sum has been set aside for {E} annually?"],
              ["Which site has an annual budget of {v}?"], lo=120, hi=980, prefix="€", suffix="k", sep=False,
              carrier="row", label="Annual budget", cmp=("higher", "lower", "annual budget")),
        _attr("lead", "person", ["{Ec} is led by site lead {v}."],
              ["Who is the site lead at {E}?", "Who leads {E}?"],
              ["Which person is responsible for {E}?", "Who is the contact in charge of {E}?"],
              ["Which site is led by {v}?"], carrier="row", label="Site lead"),
    ],
}

# ------------------------------------------------------------------- model


@dataclass
class Fact:
    domain: str
    attr: str
    place: str            # place or model name this fact is about
    entity: str           # "the Calanford depot"
    value: str            # as written: "94.2 percent", "Orlin Vasquez", "1987"
    raw: object           # comparable value (number or text)
    sentence: str         # the fact as a sentence (also the Ask key_point)
    spec: dict
    doc: str = ""
    page: Optional[int] = None
    carrier: str = "prose"

    @property
    def anchor(self) -> str:
        """The phrase a retrieved passage must contain."""
        return self.value

    def answer_regex(self) -> str:
        return answer_regex(self.value, self.spec["kind"])


@dataclass
class Section:
    title: str
    paras: list = field(default_factory=list)
    table: Optional[list] = None
    notes: Optional[str] = None


@dataclass
class DocSpec:
    filename: str
    kind: str             # docx pdf pptx xlsx csv md txt
    title: str
    sections: list = field(default_factory=list)
    rows: Optional[list] = None   # xlsx / csv: header row first


@dataclass
class Corpus:
    docs: list
    facts: list
    absent_places: list
    place_domains: dict   # place -> set of domains it appears in


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def answer_regex(value: str, kind: str) -> str:
    """A regex a correct answer is expected to match: the surname for a person,
    the digits (with an optional thousands comma) for a number."""
    if kind == "person":
        return re.escape(value.split()[-1])
    if kind == "name":
        return re.escape(value)
    number = re.search(r"\d[\d,]*\.?\d*", value)
    digits = number.group(0).replace(",", "") if number else value
    if "." not in digits and len(digits) > 3:
        digits = re.escape(digits[:-3]) + ",?" + digits[-3:]
    else:
        digits = re.escape(digits)
    return r"(?<![\d.,])" + digits + r"(?!\d)"


def _format_number(number, spec) -> str:
    if spec["kind"] == "float":
        text = f"{number:.2f}" if spec.get("prefix") == "€" else f"{number:.1f}"
    elif spec.get("sep", True):
        text = f"{number:,}"
    else:
        text = str(number)
    return f"{spec.get('prefix', '')}{text}{spec.get('suffix', '')}{spec.get('unit', '')}"


class Builder:
    def __init__(self, seed: int) -> None:
        self.rng = random.Random(seed)
        self.used_values: dict = {}
        self.used_last: set = set()
        self.used_people: set = set()
        self.facts: list = []
        self.place_domains: dict = {}

    # -- values
    def person(self) -> str:
        while True:
            last = self.rng.choice(_LAST_A) + self.rng.choice(_LAST_B)
            if last in self.used_last:
                continue
            self.used_last.add(last)
            return f"{self.rng.choice(_FIRST)} {last}"

    def value(self, domain: str, spec: dict):
        kind = spec["kind"]
        if kind == "person":
            name = self.person()
            return name, name
        if kind == "name":
            name = self.rng.choice(spec["pool"])
            return name, name
        taken = self.used_values.setdefault((domain, spec["key"]), set())
        while True:
            if kind == "float":
                number = round(self.rng.uniform(spec["lo"], spec["hi"]), 2 if spec.get("prefix") == "€" else 1)
            else:
                step = spec.get("step", 1)
                number = self.rng.randrange(spec["lo"], spec["hi"] + 1, step)
            # Warranty-style attributes have fewer possible values than entities,
            # so repeats are allowed there; reverse questions are off for them.
            if number not in taken or spec["reverse"] is None:
                taken.add(number)
                return number, _format_number(number, spec)

    def filler(self, domain: str, n: int = 2) -> str:
        return " ".join(self.rng.sample(FILLER[domain], k=min(n, len(FILLER[domain]))))

    def facts_for(self, domain: str, place: str, entity: str) -> dict:
        made = {}
        self.place_domains.setdefault(place, set()).add(domain)
        capital = entity[0].upper() + entity[1:]
        for spec in ATTRS[domain]:
            raw, value = self.value(domain, spec)
            template = self.rng.choice(spec["facts"])
            sentence = template.format(E=entity, Ec=capital, v=value)
            fact = Fact(domain, spec["key"], place, entity, value, raw, sentence, spec,
                        carrier=spec.get("carrier", "prose"))
            made[spec["key"]] = fact
            self.facts.append(fact)
        return made


def _place(fact: Fact, doc: DocSpec, page: Optional[int] = None) -> None:
    fact.doc, fact.page = doc.filename, page


# ------------------------------------------------------------------ builders


def build_corpus(seed: int = SEED) -> Corpus:
    b = Builder(seed)
    rng = b.rng
    combos = [o + c for o in _ONSETS for c in _CODAS]
    rng.shuffle(combos)
    pool, absent = combos[:40], combos[40:70]
    docs: list = []

    def draw(n: int) -> list:
        return rng.sample(pool, n)

    for place in draw(14):                                       # depot reports -> docx
        entity = f"the {place} depot"
        f = b.facts_for("depot", place, entity)
        sections = [
            Section("Overview", [f["opened"].sentence, f["manager"].sentence, b.filler("depot"), f["staff"].sentence]),
            Section("Performance", [f["ontime"].sentence, f["pallets"].sentence, b.filler("depot")]),
            Section("Costs", [f"The table below lists cost figures for {entity}."],
                    table=[["Measure", "Value"], [ATTRS["depot"][5]["label"], f["cost"].value]]),
        ]
        doc = DocSpec(f"depot-report-{slug(place)}.docx", "docx", f"Depot report: {place}", sections)
        for fact in f.values():
            _place(fact, doc)
        docs.append(doc)

    for place in draw(12):                                       # coral surveys -> pdf, one fact group per page
        entity = f"the {place} Bay survey"
        f = b.facts_for("study", place, entity)
        sections = [
            Section("Abstract", [f["lead"].sentence, b.filler("study"), f["funder"].sentence]),
            Section("Methods", [f["sample"].sentence, b.filler("study"), f["duration"].sentence]),
            Section("Results", [f["survival"].sentence, b.filler("study", 3)]),
            Section("Discussion", [b.filler("study", 3), b.filler("study", 2)]),
        ]
        doc = DocSpec(f"{slug(place)}-bay-survey.pdf", "pdf", f"{place} Bay coral survey: final report", sections)
        pages = {"lead": 1, "funder": 1, "sample": 2, "duration": 2, "survival": 3}
        for key, fact in f.items():
            _place(fact, doc, pages[key])
        docs.append(doc)

    for place in draw(10):                                       # office handbooks -> md
        entity = f"the {place} office"
        f = b.facts_for("policy", place, entity)
        sections = [
            Section("Travel and expenses", [f["travel"].sentence, f["approver"].sentence, b.filler("policy")]),
            Section("Employment", [b.filler("policy"), f["notice"].sentence]),
            Section("Governance", [f["review"].sentence, b.filler("policy")]),
        ]
        doc = DocSpec(f"{slug(place)}-office-handbook.md", "md", f"{place} office handbook", sections)
        for fact in f.values():
            _place(fact, doc)
        docs.append(doc)

    names = rng.sample(_PRODUCT_NAMES, 12)
    for name in names:                                           # product datasheets -> pptx
        model = f"{name}-{rng.randrange(100, 999)}"
        entity = f"the {model} {rng.choice(_PRODUCT_NOUNS)}"
        f = b.facts_for("product", model, entity)
        sections = [
            Section(f"{model}", ["Product datasheet"]),
            Section("Specifications", [f["weight"].sentence, f["price"].sentence, b.filler("product")]),
            Section("Endurance", [f["battery"].sentence, b.filler("product")], notes=f["depth"].sentence),
            Section("Support", [f["warranty"].sentence, b.filler("product")]),
        ]
        doc = DocSpec(f"{slug(model)}-datasheet.pptx", "pptx", f"{model} datasheet", sections)
        for fact in f.values():
            _place(fact, doc)
        f["depth"].carrier = "notes"
        docs.append(doc)

    for place in draw(10):                                       # town chronicles -> txt
        entity = f"the town of {place}"
        f = b.facts_for("town", place, entity)
        sections = [
            Section("Chapter 1: The founding", [f["founded"].sentence, b.filler("town", 3), f["founder"].sentence,
                                                b.filler("town", 2)]),
            Section("Chapter 2: The charter", [b.filler("town", 3), f["charter"].sentence, b.filler("town", 3)]),
            Section("Chapter 3: The bridge", [f["bridge"].sentence, b.filler("town", 3), f["population"].sentence,
                                              b.filler("town", 2)]),
        ]
        doc = DocSpec(f"chronicle-of-{slug(place)}.txt", "txt", f"The chronicle of {place}", sections)
        for fact in f.values():
            _place(fact, doc)
        docs.append(doc)

    sites = draw(24)                                             # site rosters -> xlsx and csv
    header = ["Site"] + [s["label"] for s in ATTRS["roster"]]
    halves = [("site-roster-north.xlsx", "xlsx", sites[:12]), ("site-roster-south.csv", "csv", sites[12:])]
    for filename, kind, members in halves:
        rows = [header]
        doc = DocSpec(filename, kind, "Site roster", rows=rows)
        for place in members:
            f = b.facts_for("roster", place, f"the {place} site")
            rows.append([place, f["code"].value, f["budget"].value, f["lead"].value])
            for fact in f.values():
                _place(fact, doc)
        docs.append(doc)

    return Corpus(docs, b.facts, absent, b.place_domains)


# ------------------------------------------------------------------- writers


def write_docx(path: Path, doc: DocSpec) -> None:
    from docx import Document

    document = Document()
    document.add_heading(doc.title, level=0)
    for section in doc.sections:
        document.add_heading(section.title, level=1)
        for text in section.paras:
            document.add_paragraph(text)
        if section.table:
            table = document.add_table(rows=0, cols=len(section.table[0]))
            table.style = "Table Grid"
            for row in section.table:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
    document.save(path)


def write_pptx(path: Path, doc: DocSpec) -> None:
    from pptx import Presentation

    deck = Presentation()
    for index, section in enumerate(doc.sections):
        if index == 0:
            slide = deck.slides.add_slide(deck.slide_layouts[0])
            slide.shapes.title.text = f"{section.title}"
            slide.placeholders[1].text = section.paras[0]
            continue
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = section.title
        body = slide.placeholders[1].text_frame
        body.text = section.paras[0]
        for text in section.paras[1:]:
            body.add_paragraph().text = text
        if section.notes:
            slide.notes_slide.notes_text_frame.text = section.notes
    deck.save(path)


def write_pdf(path: Path, doc: DocSpec) -> None:
    import fitz

    pdf = fitz.open()
    for index, section in enumerate(doc.sections):
        page = pdf.new_page()
        y = 72.0
        if index == 0:
            page.insert_text((72, y), doc.title, fontsize=17, fontname="helv")
            y += 34
        page.insert_text((72, y), section.title, fontsize=14, fontname="helv")
        y += 26
        for text in section.paras:
            for line in textwrap.wrap(text, 76):
                page.insert_text((72, y), line, fontsize=11, fontname="helv")
                y += 15
            y += 9
        if y > 780:
            raise ValueError(f"{doc.filename} page {index + 1} overflows ({y:.0f}pt)")
    pdf.save(path)
    pdf.close()


def write_text(path: Path, doc: DocSpec) -> None:
    markdown = doc.kind == "md"
    lines = [f"# {doc.title}" if markdown else doc.title.upper(), ""]
    for section in doc.sections:
        lines += [f"## {section.title}" if markdown else section.title, ""]
        for text in section.paras:
            lines += [text, ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_sheet(path: Path, doc: DocSpec) -> None:
    if doc.kind == "csv":
        with path.open("w", newline="", encoding="utf-8") as handle:
            csv.writer(handle).writerows(doc.rows)
        return
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Roster"
    for row in doc.rows:
        sheet.append(row)
    notes = book.create_sheet("Notes")
    for sentence in FILLER["roster"]:
        notes.append([sentence])
    book.save(path)


WRITERS = {"docx": write_docx, "pptx": write_pptx, "pdf": write_pdf, "md": write_text, "txt": write_text,
           "xlsx": write_sheet, "csv": write_sheet}


def write_corpus(corpus: Corpus, out: Path) -> list:
    out.mkdir(parents=True, exist_ok=True)
    paths = []
    for doc in corpus.docs:
        path = out / doc.filename
        WRITERS[doc.kind](path, doc)
        paths.append(path)
    return paths


# ------------------------------------------------------------------- queries


def _fmt(template: str, fact: Fact) -> str:
    capital = fact.entity[0].upper() + fact.entity[1:]
    return template.format(E=fact.entity, Ec=capital, v=fact.value)


def _tags(fact: Fact, qtype: str, corpus: Corpus) -> list:
    kind = fact.doc.rsplit(".", 1)[-1]
    tags = [fact.domain, qtype, kind, fact.carrier]
    if len(corpus.place_domains.get(fact.place, ())) > 1:
        tags.append("shared-name")
    return tags


def _judged(fact: Fact) -> dict:
    entry = {"document": fact.doc, "contains": fact.anchor}
    if fact.page is not None:
        entry["page"] = fact.page
    return entry


def _spread(facts: list, total: int, rng: random.Random) -> list:
    """Pick `total` facts, keeping each domain's share of the corpus, in a
    stable shuffled order."""
    by_domain: dict = {}
    for fact in facts:
        by_domain.setdefault(fact.domain, []).append(fact)
    picked = []
    for domain in sorted(by_domain):
        members = by_domain[domain][:]
        rng.shuffle(members)
        picked += members[: round(total * len(members) / len(facts))]
    rng.shuffle(picked)
    return picked


def _question(fact: Fact, qtype: str, rng: random.Random) -> tuple:
    spec = fact.spec
    if qtype == "reverse" and not spec["reverse"]:
        qtype = "para"
    return _fmt(rng.choice(spec[qtype]), fact), qtype


def _comparisons(corpus: Corpus, per_domain: int, rng: random.Random) -> list:
    """Pairs of same-attribute facts from different entities, for 'which is
    higher, A or B?' -- the answer needs both sources."""
    pairs = []
    seen = set()
    by_key: dict = {}
    for fact in corpus.facts:
        if fact.spec.get("cmp") and isinstance(fact.raw, (int, float)):
            by_key.setdefault((fact.domain, fact.attr), []).append(fact)
    domains = sorted({d for d, _ in by_key})
    for domain in domains:
        keys = [k for k in sorted(by_key) if k[0] == domain]
        for index in range(per_domain):
            first, second = rng.sample(by_key[keys[index % len(keys)]], 2)
            if first.raw == second.raw or (first.place, second.place, first.attr) in seen:
                continue
            seen.add((first.place, second.place, first.attr))
            up, down, noun = first.spec["cmp"]
            word = rng.choice([up, down])
            winner = max(first, second, key=lambda f: f.raw) if word == up else min(first, second, key=lambda f: f.raw)
            pairs.append((word, noun, first, second, winner))
    return pairs


def build_queries(corpus: Corpus, seed: int = SEED, singles: int = 270, near_miss: int = 20) -> list:
    rng = random.Random(seed + 1)
    cases = []
    cycle = ["direct", "para", "direct", "para", "reverse"]
    for index, fact in enumerate(_spread(corpus.facts, singles, rng)):
        text, qtype = _question(fact, cycle[index % len(cycle)], rng)
        cases.append({"id": f"{fact.domain}-{fact.attr}-{slug(fact.place)}-{qtype}", "query": text,
                      "relevant": [_judged(fact)], "tags": _tags(fact, qtype, corpus)})

    for word, noun, first, second, _winner in _comparisons(corpus, 4, rng):
        cases.append({
            "id": f"{first.domain}-{first.attr}-{slug(first.place)}-vs-{slug(second.place)}",
            "query": f"Which has the {word} {noun}, {first.entity} or {second.entity}?",
            "relevant": [_judged(first), _judged(second)], "match": "all",
            "tags": [first.domain, "comparison", "multi-source"]})

    domains = sorted(ATTRS)
    for index in range(near_miss):                    # a real attribute, an entity that does not exist
        domain = domains[index % len(domains)]
        spec = rng.choice([s for s in ATTRS[domain] if s["kind"] != "person"])
        place = corpus.absent_places[index]
        noun = {"depot": "depot", "study": "Bay survey", "policy": "office", "town": "town", "roster": "site",
                "product": "logger"}[domain]
        entity = f"the {place} {noun}" if domain != "town" else f"the town of {place}"
        text = rng.choice(spec["direct"]).format(E=entity)
        cases.append({"id": f"near-miss-{domain}-{slug(place)}", "query": text, "expect_empty": True,
                      "tags": ["near-miss", domain]})
    for text in OFF_TOPIC:
        cases.append({"id": "off-topic-" + slug(text)[:32], "query": text, "expect_empty": True,
                      "tags": ["off-topic"]})
    return cases


def build_answers(corpus: Corpus, seed: int = SEED, singles: int = 72) -> list:
    rng = random.Random(seed + 2)
    cases = []
    cycle = ["direct", "para", "para", "direct", "reverse"]
    siblings: dict = {}
    for fact in corpus.facts:
        siblings.setdefault((fact.domain, fact.attr), []).append(fact)

    for index, fact in enumerate(_spread(corpus.facts, singles, rng)):
        text, qtype = _question(fact, cycle[index % len(cycle)], rng)
        case = {"id": f"{fact.domain}-{fact.attr}-{slug(fact.place)}-{qtype}", "question": text,
                "must_mention": [fact.answer_regex()], "key_points": [fact.sentence],
                "must_cite": [{k: v for k, v in _judged(fact).items() if k != "contains"}],
                "tags": _tags(fact, qtype, corpus)}
        others = [s for s in siblings[(fact.domain, fact.attr)]
                  if s.place != fact.place and s.value != fact.value and s.value not in fact.value
                  and fact.value not in s.value and s.value not in text]
        if others and fact.spec["kind"] != "name":
            case["must_not_mention"] = [rng.choice(others).answer_regex()]
        cases.append(case)

    for word, noun, first, second, winner in _comparisons(corpus, 2, rng):
        cases.append({
            "id": f"{first.domain}-{first.attr}-{slug(first.place)}-vs-{slug(second.place)}",
            "question": f"Which has the {word} {noun}, {first.entity} or {second.entity}?",
            "must_mention": [re.escape(winner.place)],
            "key_points": [first.sentence, second.sentence],
            "must_cite": [{"document": first.doc}, {"document": second.doc}],
            "tags": [first.domain, "comparison", "multi-source"]})

    domains = sorted(ATTRS)
    present = [f for f in corpus.facts if f.carrier == "prose" or f.carrier == "row"]
    for index in range(16):                           # a real entity, a fact nobody wrote down
        domain = domains[index % len(domains)]
        fact = rng.choice([f for f in present if f.domain == domain])
        text = rng.choice(ABSENT_TOPICS[domain]).format(E=fact.entity)
        cases.append({"id": f"absent-attribute-{domain}-{slug(fact.place)}-{index}", "question": text,
                      "expect_refusal": True, "tags": ["absent-attribute", domain]})
    for index in range(10):                           # an entity that does not exist
        domain = domains[index % len(domains)]
        spec = rng.choice([s for s in ATTRS[domain] if s["kind"] != "person"])
        noun = {"depot": "depot", "study": "Bay survey", "policy": "office", "town": "town", "roster": "site",
                "product": "logger"}[domain]
        place = corpus.absent_places[20 + index]
        entity = f"the {place} {noun}" if domain != "town" else f"the town of {place}"
        cases.append({"id": f"near-miss-{domain}-{slug(place)}", "question": rng.choice(spec["direct"]).format(E=entity),
                      "expect_refusal": True, "tags": ["near-miss", domain]})
    for text in OFF_TOPIC[:6]:
        cases.append({"id": "off-topic-" + slug(text)[:32], "question": text, "expect_refusal": True,
                      "tags": ["off-topic"]})
    return cases


# -------------------------------------------------------------------- verify


def _squash(text: str) -> str:
    return " ".join(text.split())


def read_back(path: Path) -> list:
    """The text of a generated file, one string per page (PDF) or one for the
    whole file -- read with the plain format libraries, not the service."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import fitz

        with fitz.open(path) as pdf:
            return [page.get_text() for page in pdf]
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts += [cell.text for row in table.rows for cell in row.cells]
        return ["\n".join(parts)]
    if suffix == ".pptx":
        from pptx import Presentation

        parts = []
        for slide in Presentation(path).slides:
            parts += [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
        return ["\n".join(parts)]
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path)
        return ["\n".join(str(cell) for sheet in book for row in sheet.iter_rows(values_only=True)
                          for cell in row if cell is not None)]
    return [path.read_text(encoding="utf-8")]


def verify(corpus: Corpus, out: Path, queries: list, answers: list) -> list:
    """Problems found, as strings. Empty means the ground truth checks out."""
    problems = []
    text = {}
    for doc in corpus.docs:
        path = out / doc.filename
        if not path.exists():
            problems.append(f"missing file {doc.filename}")
            continue
        text[doc.filename.lower()] = [_squash(page) for page in read_back(path)]

    def check(case_id, entry):
        pages = text.get(entry["document"].lower())
        if pages is None:
            problems.append(f"{case_id}: unknown document {entry['document']}")
            return
        needle = _squash(entry["contains"]) if entry.get("contains") else None
        if entry.get("page") is not None:
            if entry["page"] > len(pages):
                problems.append(f"{case_id}: {entry['document']} has no page {entry['page']}")
                return
            pages = [pages[entry["page"] - 1]]
        if needle and not any(needle in page for page in pages):
            problems.append(f"{case_id}: {entry['contains']!r} not in {entry['document']}"
                            + (f" page {entry['page']}" if entry.get("page") else ""))

    ids = set()
    for case in queries:
        if case["id"] in ids:
            problems.append(f"duplicate query id {case['id']}")
        ids.add(case["id"])
        for entry in case.get("relevant", []):
            check(case["id"], entry)
    for case in answers:
        for entry in case.get("must_cite", []):
            check(case["id"], entry)
        for pattern in case.get("must_mention", []):
            sources = [text[e["document"].lower()] for e in case.get("must_cite", []) if e["document"].lower() in text]
            if not any(re.search(pattern, page) for pages in sources for page in pages):
                problems.append(f"{case['id']}: must_mention {pattern!r} not found in the cited documents")

    everything = " ".join(page.lower() for pages in text.values() for page in pages)
    for keyword in ABSENT_KEYWORDS:
        if keyword in everything:
            problems.append(f"absent-topic keyword {keyword!r} appears in the corpus")
    return problems


# ----------------------------------------------------------------------- CLI


def write_jsonl(path: Path, cases: list, header: str) -> None:
    lines = [f"# {line}" for line in header.strip().splitlines()]
    lines += [json.dumps(case, ensure_ascii=False) for case in cases]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: Optional[list] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("out", type=Path, help="folder to write the documents into")
    parser.add_argument("--seed", type=int, default=SEED)
    here = Path(__file__).parent
    parser.add_argument("--queries", type=Path, default=here / "queries_synth.jsonl")
    parser.add_argument("--answers", type=Path, default=here / "answers_synth.jsonl")
    parser.add_argument("--no-verify", action="store_true")
    args = parser.parse_args(argv)

    corpus = build_corpus(args.seed)
    queries = build_queries(corpus, args.seed)
    answers = build_answers(corpus, args.seed)
    paths = write_corpus(corpus, args.out)
    note = (f"Generated by eval/synth_corpus.py (seed {args.seed}); regenerate rather than edit.\n"
            "Ground truth is known by construction. Splits: see evalstats.split_of.")
    write_jsonl(args.queries, queries, note)
    write_jsonl(args.answers, answers, note)
    print(f"{len(paths)} documents, {len(corpus.facts)} planted facts -> {args.out}")
    print(f"{len(queries)} retrieval cases -> {args.queries}")
    print(f"{len(answers)} answer cases    -> {args.answers}")

    if args.no_verify:
        return 0
    problems = verify(corpus, args.out, queries, answers)
    for problem in problems[:25]:
        print("  PROBLEM:", problem)
    print("ground truth verified against the rendered files" if not problems
          else f"{len(problems)} problem(s) -- the benchmark is not trustworthy until these are fixed")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
