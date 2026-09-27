#!/usr/bin/env python3
"""Generate the Office eval corpus: 24 seeded Word, Excel and PowerPoint files.

    python eval/office_corpus.py ../library/documents

The files are what ``queries_office.jsonl`` and ``answers_office.jsonl`` are
judged against. The seed is fixed, so every run writes the same text; only
the files' embedded timestamps differ. Register the folder as a library named
"Documents" (Library -> Collections in the UI, or POST /collections) so the
``collection`` cases can scope to it.

Deliberate edge cases: bullet lists and tables in Word, several sheets (one
header-only) with formulas and dates in Excel, speaker notes and loose text
boxes in PowerPoint, French/German/Japanese text, and one Word file
(quarterly-logistics-review-1.docx) ending in a 6,000-character base64 blob
that must neither break chunking nor get the document dropped as corrupt.

Needs python-docx, openpyxl and python-pptx (all in requirements.txt).
"""

import argparse
import base64
import datetime as dt
import random
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

rng = random.Random(20260927)
_parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
_parser.add_argument("out", type=Path, help="folder to write into, e.g. ../library/documents")
OUT = _parser.parse_args().out
OUT.mkdir(parents=True, exist_ok=True)

TOPICS = {
    "Coral reef restoration": [
        "Coral gardening grows fragments on underwater nurseries before outplanting them to degraded reefs.",
        "Bleaching occurs when warm water drives corals to expel their symbiotic zooxanthellae algae.",
        "Survival of outplanted Acropora fragments averaged 71 percent after eighteen months at the northern site.",
        "Larval reseeding captures spawn slicks and settles larvae onto ceramic substrates in tanks.",
        "Parrotfish grazing keeps macroalgae in check and gives new coral recruits room to settle.",
    ],
    "Sourdough baking": [
        "A sourdough starter is a culture of wild yeast and lactobacilli maintained with flour and water.",
        "Bulk fermentation at 24 degrees Celsius usually takes four to six hours for a 20 percent levain dough.",
        "Autolyse rests flour and water before salt is added so gluten begins forming without kneading.",
        "Higher hydration gives a more open crumb but makes shaping harder for beginners.",
        "A cold retard overnight in the fridge deepens the sour flavour and makes scoring easier.",
    ],
    "Quarterly logistics review": [
        "On-time delivery improved to 94.2 percent after the Rotterdam cross-dock went live in July.",
        "Freight costs per pallet rose 6 percent, driven mainly by diesel surcharges on the Lyon lane.",
        "Warehouse pick accuracy reached 99.7 percent following the barcode rescan policy.",
        "Carrier Nordline missed its service-level agreement in two of three months and is under review.",
        "Inventory turns fell to 7.8 because of seasonal stock built ahead of the holiday peak.",
    ],
    "Medieval astronomy": [
        "The astrolabe let medieval astronomers compute the altitude of stars and tell time at night.",
        "Al-Battani refined the length of the solar year to 365 days, 5 hours and 46 minutes.",
        "The Alfonsine Tables, compiled in Toledo, were the standard planetary tables for three centuries.",
        "Epicycles allowed the Ptolemaic model to reproduce the retrograde motion of the planets.",
        "Monastic computus determined the date of Easter from lunar and solar cycles.",
    ],
    "Home solar installation": [
        "A 6 kilowatt rooftop array in a sunny climate produces roughly 9,000 kilowatt-hours per year.",
        "Microinverters let each panel operate independently so shading one panel does not drag down the string.",
        "Net metering credits exported electricity against the household's grid consumption.",
        "Battery storage with a 13.5 kWh pack can carry essential loads through an overnight outage.",
        "Panel degradation is typically around half a percent of output per year.",
    ],
    "Onboarding handbook": [
        "New hires receive laptops on day one and complete security training within their first week.",
        "Each newcomer is paired with an onboarding buddy from a different team for the first month.",
        "Expense reports are submitted through the finance portal and approved by the direct manager.",
        "Remote employees may claim a one-time home office stipend of 500 euros.",
        "Performance check-ins happen at 30, 60 and 90 days after the start date.",
    ],
    "Café équipe — notes multilingues": [
        "Le café de l'équipe ouvre à huit heures; la réunion hebdomadaire a lieu le mardi.",
        "Das Team trifft sich donnerstags zur Retrospektive im Konferenzraum Süd.",
        "東京オフィスの会議は毎週水曜日に行われます。",
        "Naïve café résumé: coördinate the piñata order before the jubilee.",
    ],
}

TOPIC_NAMES = list(TOPICS)


def paragraphs(topic, n):
    sentences = TOPICS[topic]
    return [" ".join(rng.sample(sentences, k=min(len(sentences), rng.randint(2, 4)))) for _ in range(n)]


def slug(text):
    return "".join(c.lower() if c.isalnum() else "-" for c in text).strip("-").replace("--", "-")


def make_docx(topic, index, blob=False):
    doc = Document()
    doc.add_heading(topic, level=0)
    doc.add_paragraph(f"Prepared {dt.date(2026, rng.randint(1, 9), rng.randint(1, 28)):%d %B %Y}.")
    for section in range(rng.randint(2, 4)):
        doc.add_heading(f"Section {section + 1}", level=1)
        for text in paragraphs(topic, rng.randint(1, 3)):
            doc.add_paragraph(text)
        if rng.random() < 0.6:
            for text in rng.sample(TOPICS[topic], k=3):
                doc.add_paragraph(text, style="List Bullet")
        if rng.random() < 0.5:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, label in zip(table.rows[0].cells, ("Item", "Value", "Note")):
                cell.text = label
            for row in range(rng.randint(2, 5)):
                cells = table.add_row().cells
                cells[0].text = f"{topic.split()[0]} item {row + 1}"
                cells[1].text = f"{rng.uniform(1, 500):.1f}"
                cells[2].text = rng.choice(["on track", "delayed", "needs review", "complete"])
    if blob:
        doc.add_heading("Appendix: embedded payload", level=1)
        payload = base64.b64encode(rng.randbytes(4500)).decode()
        run = doc.add_paragraph().add_run(payload)
        run.font.size = Pt(6)
    path = OUT / f"{slug(topic)}-{index}.docx"
    doc.save(path)
    return path


def make_xlsx(topic, index, empty_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Q1", "Q2", "Q3", "Total"])
    for row in range(rng.randint(4, 12)):
        values = [round(rng.uniform(10, 1000), 2) for _ in range(3)]
        r = row + 2
        ws.append([f"{topic.split()[0]} metric {row + 1}", *values, f"=SUM(B{r}:D{r})"])
    detail = wb.create_sheet("Detail")
    detail.append(["Date", "Owner", "Description", "Amount"])
    for _ in range(rng.randint(8, 25)):
        detail.append([
            dt.date(2026, rng.randint(1, 9), rng.randint(1, 28)),
            rng.choice(["Ana", "Bo", "Chidi", "Dana", "Eli", "Farah"]),
            rng.choice(TOPICS[topic]),
            round(rng.uniform(-200, 2000), 2),
        ])
    if empty_sheet:
        wb.create_sheet("Blank").append(["Header only"])
    path = OUT / f"{slug(topic)}-data-{index}.xlsx"
    wb.save(path)
    return path


def make_pptx(topic, index):
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = topic
    title.placeholders[1].text = f"Briefing deck {index}"
    for number in range(rng.randint(3, 7)):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"{topic}: point {number + 1}"
        body = slide.placeholders[1].text_frame
        bullets = rng.sample(TOPICS[topic], k=min(3, len(TOPICS[topic])))
        body.text = bullets[0]
        for text in bullets[1:]:
            body.add_paragraph().text = text
        if rng.random() < 0.5:
            slide.notes_slide.notes_text_frame.text = "Speaker notes: " + rng.choice(TOPICS[topic])
        if rng.random() < 0.3:
            box = slide.shapes.add_textbox(Inches(1), Inches(6.3), Inches(8), Inches(0.6))
            box.text_frame.text = "Source: internal figures, " + str(rng.randint(2019, 2026))
    path = OUT / f"{slug(topic)}-deck-{index}.pptx"
    prs.save(path)
    return path


made = []
for i, topic in enumerate(TOPIC_NAMES):
    made.append(make_docx(topic, 1, blob=(i == 2)))
    made.append(make_pptx(topic, 1))
    if i % 2 == 0:
        made.append(make_xlsx(topic, 1, empty_sheet=(i == 0)))
for topic in rng.sample(TOPIC_NAMES, 3):
    made.append(make_docx(topic, 2))
    made.append(make_xlsx(topic, 2))
for path in made:
    print(f"{path.stat().st_size:>8}  {path.name}")
print(len(made), "files")
